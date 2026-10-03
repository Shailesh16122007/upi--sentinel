"""Document and structured-file extraction: TXT, PDF (text layer + OCR for scanned pages), DOCX, CSV, JSON."""
from __future__ import annotations

import csv
import io
import json
from dataclasses import dataclass, field

from app.services.ocr import engine as ocr


@dataclass
class DocResult:
    ok: bool
    text: str = ""
    method: str = ""
    pages: int | None = None
    page_texts: list[str] = field(default_factory=list)
    quality: float | None = None
    error: str | None = None
    rows: list[dict] | None = None  # structured files


def decode_text(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-16", "cp1252", "latin-1"):
        try:
            t = data.decode(enc)
            if enc == "utf-16" and "\x00" in t:
                continue
            return t
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def extract_pdf(data: bytes) -> DocResult:
    try:
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception:  # noqa: BLE001
                return DocResult(ok=False, error="PDF is password-protected.")
        texts = [(p.extract_text() or "").strip() for p in reader.pages]
    except Exception as e:  # noqa: BLE001
        return DocResult(ok=False, error=f"Could not read PDF ({type(e).__name__}).")

    ocr_pages, quals = [], []
    scanned = [i for i, t in enumerate(texts) if len(t) < 25]
    if scanned:
        if not ocr.available():
            if len(scanned) == len(texts):
                return DocResult(ok=False, pages=len(texts), error="Scanned PDF and OCR is not available on this server.")
        else:
            try:
                import pypdfium2 as pdfium
                doc = pdfium.PdfDocument(data)
                for i in scanned[:20]:  # cap OCR work per document
                    img = doc[i].render(scale=2.5).to_pil()
                    buf = io.BytesIO()
                    img.save(buf, format="PNG")
                    r = ocr.ocr_image(buf.getvalue())
                    if r.ok:
                        texts[i] = r.text
                        ocr_pages.append(i + 1)
                        quals.append(r.quality)
            except Exception as e:  # noqa: BLE001
                if len(scanned) == len(texts):
                    return DocResult(ok=False, pages=len(texts), error=f"Scanned PDF could not be OCR'd ({type(e).__name__}).")
    full = "\n\n".join(f"[Page {i + 1}]\n{t}" for i, t in enumerate(texts) if t)
    if not full.strip():
        return DocResult(ok=False, pages=len(texts), error="No readable text found in this PDF.")
    method = "PDF_OCR" if ocr_pages and len(ocr_pages) == len(texts) else "PDF_TEXT+OCR" if ocr_pages else "PDF_TEXT"
    return DocResult(ok=True, text=full, method=method, pages=len(texts), page_texts=texts,
                     quality=round(sum(quals) / len(quals), 3) if quals else 1.0)


def extract_docx(data: bytes) -> DocResult:
    try:
        import docx
        d = docx.Document(io.BytesIO(data))
        parts = [p.text for p in d.paragraphs if p.text.strip()]
        for t in d.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text.strip() for c in row.cells))
        text = "\n".join(parts)
        if not text.strip():
            return DocResult(ok=False, error="The document contains no text.")
        return DocResult(ok=True, text=text, method="DOCX", quality=1.0)
    except Exception as e:  # noqa: BLE001
        return DocResult(ok=False, error=f"Could not read DOCX ({type(e).__name__}).")


def extract_structured(data: bytes, ext: str) -> DocResult:
    text = decode_text(data)
    rows: list[dict] = []
    try:
        if ext == "json":
            obj = json.loads(text)
            if isinstance(obj, dict):
                for k in ("transactions", "events", "logs", "records", "data", "items"):
                    if isinstance(obj.get(k), list):
                        for r in obj[k]:
                            if isinstance(r, dict):
                                rows.append({**r, "_collection": k})
                if not rows:
                    rows = [obj]
            elif isinstance(obj, list):
                rows = [r for r in obj if isinstance(r, dict)]
        else:
            sample = text[:4096]
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
            except csv.Error:
                dialect = csv.excel
            rows = [dict(r) for r in csv.DictReader(io.StringIO(text), dialect=dialect)]
    except (json.JSONDecodeError, csv.Error) as e:
        return DocResult(ok=False, error=f"Could not parse {ext.upper()} ({type(e).__name__}).")
    rows = [{(k or "").strip(): (v.strip() if isinstance(v, str) else v) for k, v in r.items()} for r in rows][:5000]
    if not rows:
        return DocResult(ok=False, error=f"No records found in the {ext.upper()} file.")
    return DocResult(ok=True, text=text[:200_000], method=f"STRUCTURED_{ext.upper()}", rows=rows, quality=1.0)
